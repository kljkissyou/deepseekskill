# -*- coding: utf-8 -*-
import docx
from collections import Counter
path = r"E:\Deepseek\妖市夜巡\需要核对的内容\《妖市夜巡》分镜脚本_导演3.0预设.docx"
d = docx.Document(path)
counts = Counter()
examples = {}
for para in d.paragraphs:
    for run in para.runs:
        rgb = run.font.color.rgb
        if rgb is not None:
            key = str(rgb)
            counts[key] += 1
            if key not in examples:
                examples[key] = run.text[:50]
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    rgb = run.font.color.rgb
                    if rgb is not None:
                        key = str(rgb)
                        counts[key] += 1
                        if key not in examples:
                            examples[key] = run.text[:50]
print("color counts:", dict(counts))
print("examples:", examples)
print("tables:", len(d.tables))
