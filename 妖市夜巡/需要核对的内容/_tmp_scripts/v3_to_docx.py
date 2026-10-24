# -*- coding: utf-8 -*-
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING

md_path = r"E:\Deepseek\妖市夜巡\需要核对的内容\《妖市夜巡》分镜脚本_导演3.0预设.md"
docx_path = r"E:\Deepseek\妖市夜巡\需要核对的内容\《妖市夜巡》分镜脚本_导演3.0预设.docx"

with open(md_path, encoding="utf-8") as f:
    lines = f.read().splitlines()

RED = RGBColor(0xC0, 0x00, 0x00)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GREEN = RGBColor(0x00, 0x64, 0x00)
PURPLE = RGBColor(0x80, 0x00, 0x80)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(10.5)
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
style.paragraph_format.line_spacing = 1.15
style.paragraph_format.space_after = Pt(4)

token_re = re.compile(r'==(.*?)==|\{\{(.*?)\}\}|%%(.*?)%%|\*\*(.*?)\*\*', re.S)

def add_tokenized(paragraph, text, default_bold=False, default_color=None):
    pos = 0
    for m in token_re.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.bold = default_bold
            if default_color is not None:
                r.font.color.rgb = default_color
        if m.group(1) is not None:  # red star marker
            r = paragraph.add_run(m.group(1))
            r.bold = True
            r.font.color.rgb = RED
        elif m.group(2) is not None:  # orange content
            r = paragraph.add_run(m.group(2))
            r.bold = True
            r.font.color.rgb = ORANGE
        elif m.group(3) is not None:  # green purpose
            r = paragraph.add_run(m.group(3))
            r.bold = True
            r.font.color.rgb = GREEN
        elif m.group(4) is not None:  # bold black
            r = paragraph.add_run(m.group(4))
            r.bold = True
            r.font.color.rgb = default_color if default_color is not None else RGBColor(0,0,0)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.bold = default_bold
        if default_color is not None:
            r.font.color.rgb = default_color

def is_table_sep(line):
    s = line.strip()
    return s.startswith('|')

def make_table(rows):
    # first row header
    cells = [c.strip() for c in rows[0].strip().strip('|').split('|')]
    # find if second row is separator; skip if starts/ends dashes
    data = []
    for row in rows[1:]:
        cs = [c.strip() for c in row.strip().strip('|').split('|')]
        if all(re.match(r'^:?-+:?$', c) for c in cs):
            continue
        data.append(cs)
    t = doc.add_table(rows=1, cols=len(cells))
    t.style = 'Light Grid Accent 1'
    for j, c in enumerate(cells):
        cell = t.rows[0].cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        add_tokenized(p, c, default_bold=True)
    for row in data:
        cells_row = t.add_row().cells
        for j in range(len(cells)):
            if j < len(row):
                p = cells_row[j].paragraphs[0]
                p.text = ''
                add_tokenized(p, row[j])
            else:
                cells_row[j].text = ''
    doc.add_paragraph()

def add_heading(text, level):
    # parse markdown heading text without leading #'s
    p = doc.add_heading('', level=level)
    add_tokenized(p, text)
    return p

i = 0
while i < len(lines):
    raw = lines[i]
    st = raw.strip()
    if not st:
        i += 1
        continue
    # table block
    if st.startswith('|'):
        block = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i])
            i += 1
        make_table(block)
        continue
    if st.startswith('# '):
        add_heading(st[2:].strip(), 1)
        i += 1
        continue
    if st.startswith('## '):
        add_heading(st[3:].strip(), 2)
        i += 1
        continue
    if st.startswith('### '):
        add_heading(st[4:].strip(), 3)
        i += 1
        continue
    if st.startswith('#### '):
        add_heading(st[5:].strip(), 4)
        i += 1
        continue
    if st.startswith('> '):
        p = doc.add_paragraph()
        add_tokenized(p, st[2:].strip())
        # make it quote style-ish
        p.paragraph_format.left_indent = Pt(12)
        i += 1
        continue
    # List bullet
    if st.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        add_tokenized(p, st[2:].strip())
        i += 1
        continue
    # special dialogue line
    dialogue_prefix = '**角色台词**：'
    if st.startswith(dialogue_prefix):
        p = doc.add_paragraph()
        r = p.add_run('角色台词：')
        r.bold = True
        content = st[len(dialogue_prefix):].strip()
        # determine if actual dialogue (not "无" or "无(...)")
        is_dialogue = not (content.startswith('无') or content.startswith('无（'))
        if is_dialogue:
            r2 = p.add_run(content)
            r2.bold = True
            r2.font.color.rgb = PURPLE
        else:
            add_tokenized(p, content)
        i += 1
        continue
    p = doc.add_paragraph()
    add_tokenized(p, st)
    i += 1

doc.save(docx_path)
print("Saved:", docx_path)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
