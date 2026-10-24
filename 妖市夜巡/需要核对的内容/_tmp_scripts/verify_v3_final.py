# -*- coding: utf-8 -*-
import docx
from collections import Counter
p = r"E:\Deepseek\妖市夜巡\需要核对的内容\《妖市夜巡》分镜脚本_导演3.0预设.docx"
d = docx.Document(p)
counts = Counter()
for para in d.paragraphs:
    for run in para.runs:
        if run.font.color and run.font.color.rgb:
            counts[str(run.font.color.rgb)] += 1
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if run.font.color and run.font.color.rgb:
                        counts[str(run.font.color.rgb)] += 1
print("tables", len(d.tables))
print("color counts", dict(counts))
print("last paragraphs:", [p.text for p in d.paragraphs[-6:]])
print("review table rows", len(d.tables[-1].rows) if d.tables else 0)
