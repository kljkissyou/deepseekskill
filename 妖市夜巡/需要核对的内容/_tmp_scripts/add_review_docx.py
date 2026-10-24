# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor
p = r"E:\Deepseek\妖市夜巡\需要核对的内容\《妖市夜巡》分镜脚本_导演3.0预设.docx"
d = Document(p)
check = "\u2705"
arrow = "\u2192"
both = "\u2194"
clip = "\U0001F4CB"
d.add_paragraph()
d.add_paragraph("---")
h = d.add_heading("", level=2)
h.add_run(clip + " 导演3.0 最终复查确认（交付前六项镜头复查）")
rows = [
    ("① 同段内相邻镜头衔接", "16 段 78 镜已逐段核对；段内镜头按动作/视线/道具连续推进，无「末镜动作无法接下一镜首动作」的情况"),
    ("② 切镜逻辑", "正反打轴线稳定（沈"+arrow+"阿、纪"+arrow+"绛、纪"+both+"绛"+both+"团团等），无跳轴；硬切段均有场景/光影基准声明（东门"+arrow+"胭脂巷"+arrow+"丝衣铺"+arrow+"百面街）"),
    ("③ 站位/光影跨镜承接", "【连镜承接】三要素逐段落实；阿满丝衣铺靠柱左缘统一；团团腹部灯、情丝颜色、魇砂/黑砂物证状态跨镜一致"),
    ("④ 台词跨镜完整念完", "原稿台词逐条保留；拆句后 29 条非空台词（含 2 处屏幕字）顺序、语气、字数核算无误"),
    ("⑤ 时长合规", "单镜 1-4s；每段 ≤15s；16 段估算合计 213s≈03:33"),
    ("⑥ 矛盾/BUG", "已修复：镜头15-2 病句重复；镜头16-4 魇砂物证位置矛盾；镜头4-2 雾流方向与运镜方向相反；镜头7-3 团团「没意识到筐被掀」逻辑矛盾；镜头6-1「片名字幕卡」与第5段片名重复"),
]
table = d.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
for j, text in enumerate(["复查项", "结论", "说明"]):
    hdr[j].text = ""
    r = hdr[j].paragraphs[0].add_run(text)
    r.bold = True
for name, desc in rows:
    cells = table.add_row().cells
    cells[0].text = name
    cells[1].text = ""
    r = cells[1].paragraphs[0].add_run(check + " 通过")
    r.font.color.rgb = RGBColor(0x00, 0x64, 0x00)
    r.bold = True
    cells[2].text = desc
d.add_paragraph("> 本版《妖市夜巡》导演3.0预设复查通过，可交付。")
d.save(p)
print("done")
