# -*- coding: utf-8 -*-
"""读取指定 docx 的结构与文本（只读，不改动源文件）。"""
import sys
from docx import Document

path = sys.argv[1]
if len(sys.argv) > 2:
    sys.stdout = open(sys.argv[2], "w", encoding="utf-8")
d = Document(path)
paras = [p for p in d.paragraphs]
nonempty = [p for p in paras if p.text.strip()]
print("== 结构 ==")
print("段落总数:", len(paras), " 非空段落:", len(nonempty))
print("表格数:", len(d.tables))
for i, t in enumerate(d.tables, 1):
    rows = len(t.rows)
    cols = len(t.columns)
    head = " | ".join(c.text.strip().replace("\n", " ") for c in t.rows[0].cells) if rows else ""
    print(f"  表{i}: {rows}行 x {cols}列  表头: {head[:120]}")
print("节数:", len(d.sections))

print("\n== 样式名分布 ==")
from collections import Counter
c = Counter(p.style.name if p.style is not None else "?" for p in nonempty)
for k, v in c.most_common():
    print(f"  {k}: {v}")

print("\n== 前 60 个非空段落 ==")
for i, p in enumerate(nonempty[:60], 1):
    txt = p.text.strip()
    print(f"[{i}][{p.style.name}] {txt[:160]}")
